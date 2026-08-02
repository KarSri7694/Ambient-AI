"""Sandboxed MCP tools for reading, creating, and editing local documents."""

from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any

import fitz
from docx import Document
from fastmcp import FastMCP
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from config import CONFIG
from local_control.filesystem import FileGrantSet
from local_control.safety import PathGrantError


mcp = FastMCP("Document Tools")
_SUPPORTED_READ_EXTENSIONS = {".pdf", ".docx", ".pptx", ".md", ".txt"}
_SUPPORTED_CREATE_FORMATS = {"pdf", "docx", "pptx", "md", "txt"}


class DocumentToolError(ValueError):
    pass


def _configured_roots() -> list[str]:
    raw = CONFIG.get_str("documents", "allowed_paths_json", "[]")
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise DocumentToolError("[documents] allowed_paths_json must be a JSON array of absolute paths.") from exc
    if not isinstance(parsed, list) or not all(isinstance(item, str) and item.strip() for item in parsed):
        raise DocumentToolError("[documents] allowed_paths_json must contain one or more absolute path strings.")
    return [str(item) for item in parsed]


def _resolve(path: str, *, must_exist: bool = False) -> Path:
    try:
        target = FileGrantSet(_configured_roots()).resolve(path)
    except PathGrantError as exc:
        raise DocumentToolError(str(exc)) from exc
    if must_exist and not target.exists():
        raise DocumentToolError(f"Path does not exist: {target}")
    return target


def _require_supported(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in _SUPPORTED_READ_EXTENSIONS:
        raise DocumentToolError(f"Unsupported type '{extension or '[no extension]'}'.")
    return extension


def _adjacent_backup(source: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = source.with_name(f"{source.stem}.{stamp}.bak{source.suffix}")
    suffix = 2
    while backup.exists():
        backup = source.with_name(f"{source.stem}.{stamp}.{suffix}.bak{source.suffix}")
        suffix += 1
    shutil.copy2(source, backup)
    return backup


def _append_markdown(document: Document, content: str) -> None:
    for raw_line in str(content or "").splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)


def _add_slide(presentation: Presentation, payload: dict[str, Any]) -> None:
    bullets = payload.get("bullets") or []
    if not isinstance(bullets, list):
        raise DocumentToolError("Each slide's bullets must be a list.")
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = str(payload.get("title") or "Untitled slide")
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for index, item in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = str(item)
    notes = str(payload.get("notes") or "").strip()
    if notes and getattr(slide, "notes_slide", None) is not None:
        notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
        if notes_frame is not None:
            notes_frame.text = notes


def _create_pdf(output: Path, title: str, content: str) -> None:
    width, height = A4
    pdf = canvas.Canvas(str(output), pagesize=A4)
    y = height - 56
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(48, y, title or output.stem)
    y -= 32
    pdf.setFont("Helvetica", 10)
    for paragraph in str(content or "").splitlines() or [""]:
        for line in re.findall(r".{1,95}(?:\s+|$)|\S+?(?:\s+|$)", paragraph) or [""]:
            if y < 54:
                pdf.showPage()
                y = height - 56
                pdf.setFont("Helvetica", 10)
            pdf.drawString(48, y, line.strip())
            y -= 14
    pdf.save()


@mcp.tool
def document_inspect(path: Annotated[str, "Absolute allowed document path"]) -> dict[str, Any]:
    """Inspect a supported document without changing it."""
    try:
        target, extension = _resolve(path, must_exist=True), ""
        extension = _require_supported(target)
        result: dict[str, Any] = {"ok": True, "path": str(target), "format": extension[1:], "size_bytes": target.stat().st_size, "editable": extension in {".pdf", ".docx", ".pptx"}}
        if extension == ".pdf":
            with fitz.open(target) as pdf:
                result.update({"page_count": pdf.page_count, "metadata": dict(pdf.metadata or {})})
        elif extension == ".docx":
            doc = Document(str(target)); result.update({"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)})
        elif extension == ".pptx":
            result["slide_count"] = len(Presentation(str(target)).slides)
        return result
    except (DocumentToolError, OSError, RuntimeError) as exc:
        return {"ok": False, "error": str(exc)}


@mcp.tool
def document_read(path: Annotated[str, "Absolute allowed PDF, DOCX, PPTX, Markdown, or text path"]) -> dict[str, Any]:
    """Read the structured content of a supported local document."""
    try:
        target = _resolve(path, must_exist=True)
        extension = _require_supported(target)
        if extension == ".pdf":
            with fitz.open(target) as pdf:
                return {"ok": True, "path": str(target), "format": "pdf", "pages": [{"page": index + 1, "text": page.get_text("text")} for index, page in enumerate(pdf)]}
        if extension == ".docx":
            doc = Document(str(target))
            return {"ok": True, "path": str(target), "format": "docx", "paragraphs": [item.text for item in doc.paragraphs if item.text.strip()], "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]}
        if extension == ".pptx":
            presentation = Presentation(str(target))
            return {"ok": True, "path": str(target), "format": "pptx", "slides": [{"slide": index + 1, "text": [str(shape.text).strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and str(shape.text).strip()]} for index, slide in enumerate(presentation.slides)]}
        return {"ok": True, "path": str(target), "format": extension[1:], "content": target.read_text(encoding="utf-8", errors="replace")}
    except (DocumentToolError, OSError, RuntimeError) as exc:
        return {"ok": False, "error": str(exc)}


@mcp.tool
def document_create(
    format: Annotated[str, "pdf, docx, pptx, md, or txt"],
    output_path: Annotated[str, "New absolute output path inside an allowed root"],
    title: Annotated[str, "Document title"] = "",
    content: Annotated[str, "Markdown-like content for PDF, DOCX, Markdown, or text"] = "",
    slides: Annotated[list[dict[str, Any]], "PPTX slide objects: title, bullets, optional notes"] = [],
) -> dict[str, Any]:
    """Create a new document without overwriting an existing file."""
    try:
        kind = str(format or "").strip().lower().lstrip(".")
        if kind not in _SUPPORTED_CREATE_FORMATS:
            raise DocumentToolError("format must be pdf, docx, pptx, md, or txt.")
        output = _resolve(output_path)
        if output.exists():
            raise DocumentToolError("Output already exists. Use document_edit for an existing file.")
        if output.suffix.lower() != f".{kind}" or not output.parent.is_dir():
            raise DocumentToolError(f"output_path must end in .{kind} and its directory must exist.")
        if kind in {"md", "txt"}:
            output.write_text(content, encoding="utf-8")
        elif kind == "docx":
            doc = Document()
            if title.strip(): doc.add_heading(title.strip(), level=0)
            _append_markdown(doc, content); doc.save(str(output))
        elif kind == "pptx":
            if not slides: raise DocumentToolError("PPTX creation requires at least one slide.")
            presentation = Presentation()
            for slide in slides:
                if not isinstance(slide, dict): raise DocumentToolError("slides entries must be objects.")
                _add_slide(presentation, slide)
            presentation.save(str(output))
        else:
            _create_pdf(output, title, content)
        return {"ok": True, "created_path": str(output), "format": kind}
    except (DocumentToolError, OSError, RuntimeError) as exc:
        return {"ok": False, "error": str(exc)}


def _edit_docx(source: Path, operations: list[dict[str, Any]], output: Path) -> None:
    doc = Document(str(source))
    for operation in operations:
        kind = str(operation.get("type") or "")
        if kind == "replace_text":
            old, new = str(operation.get("old") or ""), str(operation.get("new") or "")
            if not old: raise DocumentToolError("DOCX replace_text requires old.")
            for paragraph in list(doc.paragraphs) + [paragraph for table in doc.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]:
                if old in paragraph.text: paragraph.text = paragraph.text.replace(old, new)
        elif kind == "append_content": _append_markdown(doc, str(operation.get("content") or ""))
        else: raise DocumentToolError(f"Unsupported DOCX operation: {kind}")
    doc.save(str(output))


def _edit_pptx(source: Path, operations: list[dict[str, Any]], output: Path) -> None:
    presentation = Presentation(str(source))
    for operation in operations:
        kind = str(operation.get("type") or "")
        if kind == "replace_text":
            old, new = str(operation.get("old") or ""), str(operation.get("new") or "")
            if not old: raise DocumentToolError("PPTX replace_text requires old.")
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and old in shape.text: shape.text = shape.text.replace(old, new)
        elif kind == "append_slides":
            for slide in operation.get("slides") or []:
                if not isinstance(slide, dict): raise DocumentToolError("append_slides entries must be objects.")
                _add_slide(presentation, slide)
        else: raise DocumentToolError(f"Unsupported PPTX operation: {kind}")
    presentation.save(str(output))


def _edit_pdf(source: Path, operations: list[dict[str, Any]], output: Path) -> None:
    document = fitz.open(source)
    try:
        for operation in operations:
            kind = str(operation.get("type") or "")
            if kind in {"reorder_pages", "extract_pages"}:
                pages = operation.get("pages") or []
                if not isinstance(pages, list) or not pages: raise DocumentToolError(f"PDF {kind} requires pages.")
                replacement = fitz.open()
                for number in pages:
                    index = int(number) - 1
                    if index < 0 or index >= document.page_count: raise DocumentToolError(f"PDF page {number} is outside the document.")
                    replacement.insert_pdf(document, from_page=index, to_page=index)
                document.close(); document = replacement
            elif kind == "merge_pdf":
                other = _resolve(str(operation.get("path") or ""), must_exist=True)
                if other.suffix.lower() != ".pdf": raise DocumentToolError("merge_pdf requires a PDF path.")
                with fitz.open(other) as extra: document.insert_pdf(extra)
            elif kind == "watermark_text":
                text = str(operation.get("text") or "").strip()
                if not text: raise DocumentToolError("watermark_text requires text.")
                for page in document: page.insert_text((36, 36), text, fontsize=10, color=(0.45, 0.45, 0.45), overlay=True)
            elif kind == "add_text_annotation":
                page_number, text = int(operation.get("page") or 1), str(operation.get("text") or "").strip()
                if not text or page_number < 1 or page_number > document.page_count: raise DocumentToolError("Invalid add_text_annotation operation.")
                document[page_number - 1].add_text_annot((float(operation.get("x") or 36), float(operation.get("y") or 36)), text)
            else: raise DocumentToolError(f"Unsupported PDF operation: {kind}")
        document.save(str(output))
    finally:
        document.close()


@mcp.tool
def document_edit(
    path: Annotated[str, "Absolute DOCX, PPTX, or PDF path inside an allowed root"],
    operations: Annotated[list[dict[str, Any]], "Format-specific edit operations"],
    output_path: Annotated[str, "Required for PDF; optional new output for DOCX/PPTX"] = "",
) -> dict[str, Any]:
    """Edit a document safely; in-place DOCX/PPTX changes receive adjacent backups."""
    try:
        source = _resolve(path, must_exist=True); extension = _require_supported(source)
        if extension not in {".pdf", ".docx", ".pptx"}: raise DocumentToolError("Only PDF, DOCX, and PPTX support edits.")
        if not operations or not all(isinstance(item, dict) for item in operations): raise DocumentToolError("operations must be a non-empty list of objects.")
        requested = str(output_path or "").strip()
        if extension == ".pdf":
            if not requested: raise DocumentToolError("PDF edits require a separate output_path.")
            output = _resolve(requested)
            if output == source or output.exists() or output.suffix.lower() != ".pdf": raise DocumentToolError("PDF output_path must be a new .pdf path different from the source.")
            _edit_pdf(source, operations, output)
            return {"ok": True, "source_path": str(source), "edited_path": str(output), "backup_path": None, "format": "pdf"}
        output = _resolve(requested) if requested else source
        if output != source and output.exists(): raise DocumentToolError("output_path already exists.")
        if output.suffix.lower() != extension: raise DocumentToolError(f"output_path must end in {extension}.")
        backup = _adjacent_backup(source) if output == source else None
        (_edit_docx if extension == ".docx" else _edit_pptx)(source, operations, output)
        return {"ok": True, "source_path": str(source), "edited_path": str(output), "backup_path": str(backup) if backup else None, "format": extension[1:]}
    except (DocumentToolError, OSError, RuntimeError, ValueError) as exc:
        return {"ok": False, "error": str(exc)}


@mcp.tool
def document_list_backups(path: Annotated[str, "Absolute source document path inside an allowed root"]) -> dict[str, Any]:
    """List recoverable adjacent backups created before in-place edits."""
    try:
        source = _resolve(path, must_exist=True)
        backups = sorted(source.parent.glob(f"{source.stem}.*.bak{source.suffix}"), key=lambda item: item.stat().st_mtime, reverse=True)
        return {"ok": True, "source_path": str(source), "backups": [str(item) for item in backups]}
    except (DocumentToolError, OSError) as exc:
        return {"ok": False, "error": str(exc)}


if __name__ == "__main__":
    mcp.run()
